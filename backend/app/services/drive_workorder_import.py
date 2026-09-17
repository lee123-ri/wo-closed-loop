"""钉盘「工单版」xlsx → 工单导入。

数据源：钉盘「定稿」文件夹（DRIVE_PLAN_SPACE_ID/DRIVE_PLAN_FOLDER_ID，默认指向
spaceId=26687726819 fileId=230610076351）+「初稿」文件夹（DRIVE_DRAFT_PLAN_FOLDER_ID，
默认 alidocs 节点「Agent运营计划初稿（待判定会及确认工单）」，同 space）内各项目的
「工单版」表（`目标→拆解→工单` sheet）。**只扫这两个文件夹里的「工单版」文件，不做关键词回退**
（2026-09-03 拍板：关键词回退曾把「我的文件」里的初稿静默捞进库，属静默污染，已移除；
2026-09-16 起初稿文件夹并列为轮询源）。
流程：dws 列定稿/初稿文件夹 → 下载 → 按表头自适应解析（多套列结构）→ 导入工单(source=plan)。
时间列（计划时间/时间窗口）经 plan_time 解析出 计划开始+截止。
无「工单载体」列的文件按非EAM整表导入（业务拍板 2026-09-03）；重导入仅补空日期（补空不补错）。

定稿夹是 alidocs（钉钉文档/知识库）节点，与普通钉盘节点不同（终端实测 2026-09-03）：
- 列举：`doc +list --folder <alidocs uuid>`（`drive +list` 必 RESOURCE_NOT_FOUND）；
  配置里的数字 dentryId 需经 `drive +info --node <dentryId> --space-id <spaceId>` 换 uuid，
  uuid 在响应的 result 层（顶层 fileId 为 None）。
- 节点：`{nodeId: uuid, nodeType: "file", name(不带 .xlsx 后缀), url}`，无 extension 字段。
- 下载：只有 `drive download --node <数字 dentryId>`（doc 无 +download 子命令）；
  uuid 需再经 `drive +info --space-id` 换数字 dentryId。
"""
from __future__ import annotations

import os
import re
import tempfile
from datetime import date
from pathlib import Path

import openpyxl

from app.core.database import SessionLocal
from app.models import Project, StatusLog, User, WorkOrder
from app.services.dws_client import run_dws
from app.services.plan_time import parse_deadline, parse_time_window

# 定稿文件夹（用户指定 2026-09-03）；可用环境变量覆盖
PLAN_SPACE_ID = os.environ.get("DRIVE_PLAN_SPACE_ID", "26687726819")
PLAN_FOLDER_ID = os.environ.get("DRIVE_PLAN_FOLDER_ID", "230610076351")

# 初稿文件夹（alidocs 节点「Agent运营计划初稿（待判定会及确认工单）」，2026-09-16 新增轮询源）。
# 与定稿夹同 space（26687726819），结构=项目子文件夹 × 每项目 4 类文件（视觉/汇报/文档/工单版），只导工单版。
DRAFT_PLAN_FOLDER_ID_DEFAULT = "b9Y4gmKWrPqYGdg9i4y44M2AJGXn6lpz"


def _draft_folder_id() -> str:
    """初稿文件夹 id（alidocs uuid）。优先走 Settings（读 .env，不回写 os.environ），
    兜底回退环境变量 / 默认值（沙盒无 pydantic_settings 时也能跑）。"""
    try:
        from app.core.config import get_settings
        return get_settings().drive_draft_plan_folder_id or DRAFT_PLAN_FOLDER_ID_DEFAULT
    except Exception:
        return os.environ.get("DRIVE_DRAFT_PLAN_FOLDER_ID", DRAFT_PLAN_FOLDER_ID_DEFAULT)


def find_workorder_versions() -> tuple[list[dict], list[str]]:
    """搜索年度计划工单：扫「定稿」文件夹 +「初稿」文件夹，取并集（按节点 id 去重）。

    初稿夹是 2026-09-16 用户拍板新增的轮询源（原「只扫定稿、初稿当污染排除」口径反转）：
    两夹都在同一 space（26687726819），目录结构相同（项目子文件夹各含「工单版」表）。

    返回 (files, warnings)。列举失败/为空不回退关键词搜索（初稿静默污染防线，2026-09-03 拍板），
    但单个文件夹失败/为空不阻断另一个文件夹；两夹都空则返回空 files + 明示错误。
    """
    warnings: list[str] = []
    files: list[dict] = []
    seen: set[str] = set()
    sources = (
        (PLAN_FOLDER_ID, "定稿"),
        (_draft_folder_id(), "初稿"),
    )
    for folder_id, label in sources:
        if not folder_id:
            warnings.append(f"未配置{label}文件夹，已停止，不回退关键词")
            continue
        try:
            found = _walk_folder(PLAN_SPACE_ID, folder_id)
        except Exception as e:
            warnings.append(f"{label}文件夹列举失败（space={PLAN_SPACE_ID} folder={folder_id}）：{e}；已停止，不回退关键词")
            continue
        if not found:
            warnings.append(f"{label}文件夹列举为空（无 xlsx，space={PLAN_SPACE_ID} folder={folder_id}），已跳过")
            continue
        for f in found:
            fid = f.get("id") or f.get("alt") or f.get("name")
            if fid and fid in seen:
                continue
            seen.add(fid)
            files.append(f)
    return files, warnings


def _download_ids(file: dict) -> list[str]:
    """收集下载候选节点 id：节点自身 id（alidocs uuid）+ 经
    `drive +info --node <id> --space-id <spaceId>` 换来的其它形态（数字 dentryId 等）。

    实测（2026-09-03）：doc 无 +download 子命令，下载只有 `drive download --node`；
    doc +list 给的是 alidocs uuid，drive download 可能只认数字 dentryId，故逐层都收。
    """
    ids: list[str] = []
    for i in (file.get("id"), file.get("alt")):
        if i and str(i) not in ids:
            ids.append(str(i))
    for i in list(ids):
        try:
            info = run_dws("drive", "+info", "--node", i, "--space-id", PLAN_SPACE_ID)
        except Exception:
            continue
        layers = [info] if isinstance(info, dict) else []
        for k in ("result", "data"):
            v = info.get(k) if isinstance(info, dict) else None
            if isinstance(v, dict):
                layers.append(v)
        for layer in layers:
            for k in ("dentryId", "fileId", "uuid", "dentryUuid", "nodeId", "id"):
                v = layer.get(k)
                if v and str(v) not in ids:
                    ids.append(str(v))
    return ids


def _download(file: dict) -> Path:
    """下载文件：`drive download --node <id>`（唯一的下载通道），逐候选 id 尝试，落盘为准。"""
    tmp = Path(tempfile.mkdtemp(prefix="wo_drive_"))
    for i in _download_ids(file):
        try:
            # download 的 stdout 不保证是 JSON，不做解析，只看落盘文件
            run_dws("drive", "download", "--node", i, "--output", str(tmp), parse_json=False)
        except Exception:
            continue
        xs = list(tmp.glob("*.xlsx")) + list(tmp.glob("*.xls"))
        if xs:
            return xs[0]
    return tmp


def _norm_node(n: dict) -> dict | None:
    """兼容 alidocs（doc +list：nodes/nodeId/nodeType/url）与
    普通钉盘（drive +list：files/dentryId/type）两种节点结构。

    实测（2026-09-03）：定稿夹 doc +list 的节点形如
    {nodeId: uuid, nodeType: "file", name: "…年度运营计划（工单版）", url: …}
    ——无 extension 字段、文件名不带 .xlsx 后缀。故：
    - extension 存在且非 xlsx/xls → 拒；
    - extension 缺失但 nodeType=FILE → 收（是否真表格由下载+解析兜底，失败记 errors）。

    返回 {kind: file|folder, name, id, alt}；明确非表格的文件返回 None。
    """
    if not isinstance(n, dict):
        return None
    name = str(n.get("name") or n.get("title") or "")
    ntype = str(n.get("nodeType") or n.get("type") or "").upper()
    ext = str(n.get("extension") or "").lower().lstrip(".")
    ids = [str(n[k]) for k in ("fileId", "nodeId", "dentryId", "id") if n.get(k)]
    if not ids:
        return None
    head = {"id": ids[0], "alt": ids[1] if len(ids) > 1 else None, "name": name}
    if ntype in ("FOLDER", "DIR"):
        return {"kind": "folder", **head}
    if name.lower().endswith((".docx", ".doc", ".pdf", ".pptx", ".ppt", ".txt", ".md", ".zip")):
        return None
    if ext:
        if ext in ("xlsx", "xls"):
            return {"kind": "file", **head}
        return None
    if ntype in ("FILE", ""):
        return {"kind": "file", **head}
    return None


def _resolve_doc_folder(folder_id: str) -> str:
    """doc +list 的 --folder 要 alidocs uuid；用户给的可能是数字 dentryId，经
    `drive +info --node <id> --space-id <spaceId>` 换 uuid。

    实测（2026-09-03）：+info 不带 --space-id 直接报「缺少必要信息 spaceId」；
    uuid 在响应的 result 层，顶层 fileId 为 None。逐层取，取不到则原样返回
    （让上层报真实错误，不静默吞）。
    """
    if not str(folder_id).isdigit():
        return folder_id
    info = run_dws("drive", "+info", "--node", folder_id, "--space-id", PLAN_SPACE_ID)
    layers = [info]
    for k in ("result", "data"):
        v = info.get(k) if isinstance(info, dict) else None
        if isinstance(v, dict):
            layers.append(v)
    for layer in layers:
        for k in ("fileId", "uuid", "dentryUuid", "nodeId"):
            v = layer.get(k)
            if v and not str(v).isdigit():
                return str(v)
    return folder_id


def _walk_folder(space_id: str, folder_id: str, depth: int = 0) -> list[dict]:
    """递归列举定稿文件夹内的 xlsx（子文件夹最多 3 层）。

    定稿文件夹是钉钉文档/知识库(alidocs)节点，drive +list 列举不到（RESOURCE_NOT_FOUND），
    必须走 doc +list（终端实测 2026-09-03）。space_id 仅保留作环境变量兼容，doc 通道不用。
    """
    if depth > 3:
        return []
    doc_id = _resolve_doc_folder(folder_id) if depth == 0 else folder_id
    data = run_dws("doc", "+list", "--folder", doc_id)
    nodes = data.get("nodes") or data.get("items") or data.get("files") or data.get("list") or []
    out: list[dict] = []
    for n in nodes:
        norm = _norm_node(n)
        if not norm:
            continue
        if norm["kind"] == "folder":
            out += _walk_folder(space_id, norm["id"], depth + 1)
        # 初稿/定稿夹内每项目有视觉版/汇报版/文档版/工单版等多类节点，只收「工单版」
        elif "工单版" in (norm.get("name") or ""):
            out.append(norm)
    return out


def _map_priority(p: str) -> str:
    if not p:
        return "P2"
    m = re.match(r"P(\d)", str(p).strip())
    if not m:
        return "P2"
    return {"0": "P1", "1": "P2", "2": "P3"}.get(m.group(1), "P2")


def split_by_carrier(wos: list[dict]) -> tuple[list[dict], bool]:
    """按「工单载体」过滤：有该列 → 只留非EAM；无该列 → 整表按非EAM（业务拍板 2026-09-03）。

    返回 (保留的行, 是否有载体列)。
    """
    has_carrier = any((w.get("carrier") or "").strip() for w in wos)
    if not has_carrier:
        return list(wos), False
    return [w for w in wos if (w.get("carrier") or "").strip() == "非EAM"], True


def date_backfill(old_start: date | None, old_end: date | None,
                  new_start: date | None, new_end: date | None) -> dict:
    """补空不补错：只返回原工单缺失、新解析能补的日期。"""
    out: dict[str, date] = {}
    if old_start is None and new_start:
        out["start"] = new_start
    if old_end is None and new_end:
        out["end"] = new_end
    return out


def _planned_start(w: dict) -> date | None:
    """计划开始：优先独立「计划开始」列，否则从时间窗口左端取。"""
    s1, _ = parse_time_window(w.get("start"))
    if s1:
        return s1
    s2, _ = parse_time_window(w.get("plan"))
    return s2


def _detect_carrier_col(data_rows: list[tuple]) -> int | None:
    """按数值扫描 EAM/非EAM 标记列：某列非空值全部是 EAM/非EAM 即命中（≥3 个非空）。"""
    ncols = max((len(r) for r in data_rows), default=0)
    for j in range(ncols):
        vals = [str(r[j]).strip() for r in data_rows if j < len(r) and r[j] not in (None, "")]
        if len(vals) >= 3 and all(v in ("EAM", "非EAM") for v in vals):
            return j
    return None


def _parse_sheet(file_path: Path) -> list[dict]:
    """解析年度计划工单 sheet，兼容两种模板：

    A. 工单型（中建投/协合）：「工单编号 WO-xxx」+「工单载体」列，表含「目标→拆解→工单」
    B. 资管型（泰康）：「类别/序号/详细措施」结构，EAM/非EAM 标记在某个「备注」列

    载体标记列优先取表头含「载体/EAM」的列；否则按数值扫描 EAM/非EAM 列。
    """
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet = next((ws for ws in wb.worksheets if "目标" in ws.title or "工单" in ws.title or "计划" in ws.title), None) \
        or wb.worksheets[0]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]

    hdr_idx = None
    for i, row in enumerate(rows):
        if row and any(c and any(k in str(c) for k in ("工单编号", "详细措施", "工单名称", "工单内容")) for c in row):
            hdr_idx = i
            break
    if hdr_idx is None:
        return []

    header = [str(c).strip() if c is not None else "" for c in rows[hdr_idx]]

    def col(*keys: str) -> int | None:
        for k in keys:
            for i, h in enumerate(header):
                if k in h:
                    return i
        return None

    wo_code_idx = col("工单编号")  # None 表示资管型模板

    # 载体列：表头优先，否则按值扫描
    carrier_idx = col("工单载体", "EAM")
    if carrier_idx is None:
        carrier_idx = _detect_carrier_col(rows[hdr_idx + 1:])

    title_idx = col("工单名称", "工单内容", "详细措施", "措施事", "干什么")
    prio_idx = col("优先级")
    plan_idx = col("计划完成时间", "计划完成", "计划时间", "计划月份", "时间窗口")  # 结束 → deadline
    start_idx = col("计划开始时间", "计划开始")  # 开始 → planned_start
    person_idx = col("负责角色", "责任人", "责任部门")
    action_idx = col("做什么事", "怎么干")
    accept_idx = col("交付物", "验收")
    target_idx = col("预期效果", "对目标的价值", "预计带来的效果")
    reason_ids = [i for i, h in enumerate(header) if any(k in h for k in ("根因", "问题", "拆解", "类别"))]

    out: list[dict] = []
    for row in rows[hdr_idx + 1:]:
        vals = [str(c).strip() if c is not None else "" for c in row]
        if not any(vals):
            continue

        def g(i):
            return vals[i] if i is not None and i < len(vals) else ""

        code = g(wo_code_idx) if wo_code_idx is not None else ""
        title_raw = g(title_idx)
        if wo_code_idx is not None:
            if not code.startswith("WO-"):
                continue
        else:
            if not title_raw:
                continue  # 资管型：措施空则跳过（续行/空行）

        reason = "；".join(x for x in (g(i) for i in reason_ids if i is not None and i < len(vals))
                           if x and x not in ("—", "-"))
        out.append({
            "code": code,
            "title": title_raw,
            "action": g(action_idx) or title_raw,
            "priority": g(prio_idx),
            "reason": reason,
            "target": g(target_idx),
            "plan": g(plan_idx),
            "start": g(start_idx),
            "person": g(person_idx),
            "accept": g(accept_idx),
            "carrier": g(carrier_idx),
        })
    return out


def _match_project(filename: str, projects: list[Project]) -> Project | None:
    """按文件名匹配项目：项目简称是文件名的子串即命中（最长匹配优先）。"""
    hits = [p for p in projects if p.name and p.name in filename]
    return max(hits, key=lambda p: len(p.name)) if hits else None


def _match_project_by_content(file_path: Path, projects: list[Project]) -> Project | None:
    """文件名通用（如「工单版-运营计划」）匹配不到时的兜底：
    扫每个 sheet 前 3 行（R1 标题行，如「邯郸张西堡 2026-2027年度运营计划…」），
    单元格含项目名即命中（最长匹配优先）。只扫标题行，避免正文提及造成误归因。
    """
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    except Exception:
        return None
    hits: list[Project] = []
    try:
        for ws in wb.worksheets:
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= 3:
                    break
                for cell in row:
                    if cell is None:
                        continue
                    s = str(cell).strip()
                    if not s:
                        continue
                    for p in projects:
                        if p.name and len(p.name) >= 2 and p.name in s and p not in hits:
                            hits.append(p)
    finally:
        wb.close()
    return max(hits, key=lambda p: len(p.name)) if hits else None


def import_drive_workorder_versions() -> dict:
    """一键导入钉盘「工单版」：搜索→下载→解析→导入（按 项目+标题 去重）。"""
    files, warnings = find_workorder_versions()
    if not files:
        return {"imported": 0, "skipped_file": 0, "backfilled": 0,
                "errors": warnings or ["钉盘未找到「工单版」文件"], "files": 0}

    projects = _load_projects()
    db = SessionLocal()
    users = {u.name: u for u in db.query(User).all()}
    imported = 0
    skipped_file = 0
    backfilled = 0
    errors: list[str] = list(warnings)
    try:
        for f in files:
            try:
                path = _download(f)
                if path.is_dir():
                    skipped_file += 1
                    errors.append(f"下载失败/无xlsx: {f['name']}")
                    continue
            except Exception as e:
                skipped_file += 1
                errors.append(f"解析失败 {f['name']}: {e}")
                continue
            # 文件名优先；通用名匹配不到时按内容（R1 标题行）兜底归因
            project = _match_project(f["name"], projects) or _match_project_by_content(path, projects)
            if not project:
                skipped_file += 1
                errors.append(f"文件名与内容均匹配不到项目: {f['name']}")
                continue
            try:
                wos = _parse_sheet(path)
            except Exception as e:
                skipped_file += 1
                errors.append(f"解析失败 {f['name']}: {e}")
                continue

            # 载体列：有 → 只导非EAM；无 → 整表按非EAM导入（业务拍板 2026-09-03）
            if not wos:
                skipped_file += 1
                errors.append(f"无可解析工单: {f['name']}")
                continue
            wos, has_carrier = split_by_carrier(wos)
            if not wos:
                skipped_file += 1
                errors.append(f"全为EAM载体，跳过: {f['name']}")
                continue

            for w in wos:
                # 标题只留原始措施名，不把「WO-xxx」原编号塞进标题（原编号已记在 StatusLog 备注里）。
                title = (w["title"] or "").strip()[:256]
                # 去重：同项目同来源同标题已存在则跳过
                dup = db.query(WorkOrder).filter(
                    WorkOrder.title == title,
                    WorkOrder.project_id == project.id,
                    WorkOrder.source_code == "plan",
                ).first()
                if dup:
                    # 补空不补错：标题命中且原单日期为空 → 用新解析补上
                    fill = date_backfill(dup.planned_start_date, dup.deadline,
                                         _planned_start(w), parse_deadline(w["plan"] or w.get("start")))
                    if fill:
                        if "start" in fill:
                            dup.planned_start_date = fill["start"]
                        if "end" in fill:
                            dup.deadline = fill["end"]
                        parts = ("计划开始" if "start" in fill else "", "截止" if "end" in fill else "")
                        db.add(StatusLog(work_order_id=dup.id, from_status=dup.status,
                                         to_status=dup.status,
                                         note="重导入补空·" + "+".join(p for p in parts if p)))
                        backfilled += 1
                    continue
                reason = w["reason"]
                if w.get("target"):
                    reason = f"{reason}\n【预期】{w['target']}".strip()
                action = w["action"]
                if w.get("accept"):
                    action = f"{action}\n【验收】{w['accept']}"
                year = date.today().year
                cnt = db.query(WorkOrder).filter(WorkOrder.code.like(f"RW-{year}-%")).count()
                code = f"RW-{year}-{cnt + 1:04d}"
                person_user = users.get((w.get("person") or "").strip())
                wo = WorkOrder(
                    code=code, title=title, reason=reason or None, action=action or title,
                    project_id=project.id,
                    person_id=person_user.id if person_user else None,
                    source_code="plan", status="pending", priority=_map_priority(w["priority"]),
                    region=project.region,
                    created_date=date.today(),
                    deadline=parse_deadline(w["plan"] or w.get("start")),
                    planned_start_date=_planned_start(w),
                )
                db.add(wo)
                db.flush()
                note = f"年度计划工单导入(钉盘)·原编号{w['code']}"
                if not has_carrier:
                    note += "·无载体列按非EAM导入"
                if not person_user and w.get("person"):
                    note += f"·责任人待确认({w['person']})"
                db.add(StatusLog(work_order_id=wo.id, from_status=None, to_status="pending", note=note))
                imported += 1
        db.commit()
    finally:
        db.close()
    return {"imported": imported, "skipped_file": skipped_file, "backfilled": backfilled,
            "errors": errors, "files": len(files)}


def _load_projects() -> list[Project]:
    db = SessionLocal()
    try:
        return db.query(Project).all()
    finally:
        db.close()